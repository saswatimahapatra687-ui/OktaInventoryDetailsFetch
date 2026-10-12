#!/usr/bin/env python3
"""
Okta Inventory Script
Pulls all Okta tenant data via API and exports to Excel.

Usage:
    python okta_inventory.py --url https://demo-peach-guanaco-53879.okta.com --token 0015Obh8YMz197cgVucLKkOy_lna8slEFfJ1snIrsS

Requirements:
    pip install requests openpyxl pandas
"""

import requests
import argparse
import json
import sys
import time
import urllib3
from urllib.parse import urlparse, urlunparse
from concurrent.futures import ThreadPoolExecutor, as_completed

from datetime import datetime
from collections import defaultdict

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

import openpyxl
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side)
from openpyxl.utils import get_column_letter
import pandas as pd


# ─── OKTA API CLIENT ──────────────────────────────────────────────────────────

class OktaClient:
    def __init__(self, org_url, token):
        self.base = org_url.rstrip('/')
        self.timeout = 30
        self.headers = {
            "Authorization": f"SSWS {token}",
            "Accept": "application/json",
            "Content-Type": "application/json"
        }

    def _request(self, url, params=None):
        for attempt in range(3):
            try:
                r = requests.get(url, headers=self.headers, params=params,
                                 verify=False, timeout=self.timeout)
                if r.status_code >= 500 and attempt < 2:
                    print(f"  ⚠ Server error {r.status_code} — retrying in {2**attempt}s...")
                    time.sleep(2 ** attempt)
                    continue
                return r
            except requests.exceptions.ConnectionError as e:
                print(f"  ✗ Connection error: {e}")
                return None
            except requests.exceptions.Timeout:
                print(f"  ⚠ Timeout on attempt {attempt + 1} for {url}...")
                if attempt < 2:
                    time.sleep(2 ** attempt)
                else:
                    print(f"  ✗ All retries exhausted for {url}")
                    return None
        return None

    def get_all(self, endpoint, params=None):
        url = f"{self.base}{endpoint}"
        results = []
        while url:
            r = self._request(url, params)
            if r is None:
                return results
            if r.status_code == 401:
                print(f"  ✗ Auth failed — check your API token.")
                return results
            if r.status_code == 403:
                print(f"  ✗ Permission denied for {endpoint} — skipping.")
                return results
            if r.status_code == 429:
                reset_ts = r.headers.get("X-Rate-Limit-Reset")
                if reset_ts:
                    wait = max(int(reset_ts) - int(time.time()), 1)
                else:
                    wait = 30
                print(f"  ⏳ Rate limited — waiting {wait}s before retrying {endpoint}...")
                time.sleep(wait + 1)
                continue
            if r.status_code != 200:
                print(f"  ✗ Error {r.status_code} for {endpoint}")
                return results
            data = r.json()
            if isinstance(data, list):
                results.extend(data)
            else:
                results.append(data)
            # Rewrite Link header host to match configured base URL
            # (Okta sometimes returns internal hostnames in Link headers)
            links = r.headers.get("Link", "")
            url = None
            for link in links.split(","):
                if 'rel="next"' in link:
                    next_url = link.strip().split(";")[0].strip("<> ")
                    parsed_next = urlparse(next_url)
                    parsed_base = urlparse(self.base)
                    url = urlunparse(parsed_next._replace(
                        scheme=parsed_base.scheme, netloc=parsed_base.netloc))
            params = None
        return results

    def get_one(self, endpoint):
        url = f"{self.base}{endpoint}"
        r = self._request(url)
        if r is None:
            return {}
        if r.status_code == 200:
            return r.json()
        return {}

    def fetch_app_assignments_parallel(self, apps, group_map, workers=5):
        results = {}
        total = len(apps)
        completed = [0]

        def fetch_one(app):
            app_id = app['id']
            app_users = self.get_all(f"/api/v1/apps/{app_id}/users")
            app_groups = self.get_all(f"/api/v1/apps/{app_id}/groups")
            return app_id, {
                'user_count': len(app_users),
                'group_names': [group_map.get(g['id'], g['id']) for g in app_groups],
            }

        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(fetch_one, app): app for app in apps}
            for future in as_completed(futures):
                completed[0] += 1
                try:
                    app_id, assignment = future.result()
                    results[app_id] = assignment
                except Exception as e:
                    app = futures[future]
                    print(f"  ✗ Error fetching assignments for {app.get('label', app.get('id', ''))}: {e}")
                if completed[0] % 10 == 0 or completed[0] == total:
                    print(f"    {completed[0]}/{total} apps processed...")

        return results

    def fetch_policy_rules_parallel(self, policies, workers=5):
        results = {}

        def fetch_one(policy):
            pol_id = policy['id']
            rules = self.get_all(f"/api/v1/policies/{pol_id}/rules")
            return pol_id, rules

        with ThreadPoolExecutor(max_workers=workers) as executor:
            futures = {executor.submit(fetch_one, pol): pol for pol in policies}
            for future in as_completed(futures):
                try:
                    pol_id, rules = future.result()
                    results[pol_id] = rules
                except Exception as e:
                    pol = futures[future]
                    print(f"  ✗ Error fetching rules for policy {pol.get('name', pol.get('id', ''))}: {e}")

        return results


# ─── DATA COLLECTORS ──────────────────────────────────────────────────────────

def collect_all(client):
    data = {}

    print("\n📦 Collecting Okta Inventory...")

    print("  → Applications...")
    data['apps'] = client.get_all("/api/v1/apps?limit=200")

    print("  → Users...")
    data['users'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22ACTIVE%22")
    data['users_staged'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22STAGED%22")
    data['users_deprovisioned'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22DEPROVISIONED%22")
    data['users_suspended'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22SUSPENDED%22")
    data['users_locked'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22LOCKED_OUT%22")
    data['users_password_expired'] = client.get_all("/api/v1/users?limit=200&filter=status+eq+%22PASSWORD_EXPIRED%22")

    print("  → Groups...")
    data['groups'] = client.get_all("/api/v1/groups?limit=200")

    print("  → Identity Providers...")
    data['idps'] = client.get_all("/api/v1/idps")

    print("  → Directory Integrations (Agents)...")
    data['agents'] = client.get_all("/api/v1/agentPools")

    print("  → Policies...")
    data['auth_policies'] = client.get_all("/api/v1/policies?type=ACCESS_POLICY")
    data['mfa_policies'] = client.get_all("/api/v1/policies?type=MFA_ENROLL")
    data['password_policies'] = client.get_all("/api/v1/policies?type=PASSWORD")
    data['signon_policies'] = client.get_all("/api/v1/policies?type=OKTA_SIGN_ON")

    print("  → Policy Rules...")
    all_pols = (data.get('auth_policies', []) + data.get('mfa_policies', []) +
                data.get('password_policies', []) + data.get('signon_policies', []))
    data['policy_rules'] = client.fetch_policy_rules_parallel(all_pols)

    print("  → Features...")
    data['features'] = client.get_all("/api/v1/features")

    print("  → Trusted Origins...")
    data['trusted_origins'] = client.get_all("/api/v1/trustedOrigins")

    print("  → API Tokens...")
    data['tokens'] = client.get_all("/api/v1/api-tokens")

    print("  → Org Settings...")
    data['org'] = client.get_one("/api/v1/org")

    print("  → Network Zones...")
    data['zones'] = client.get_all("/api/v1/zones")

    print("  → Authorization Servers...")
    data['auth_servers'] = client.get_all("/api/v1/authorizationServers")

    print("  → Authenticators...")
    data['authenticators'] = client.get_all("/api/v1/authenticators")

    print("  → Group Rules...")
    data['group_rules'] = client.get_all("/api/v1/groups/rules?limit=200")

    print("  → Event Hooks...")
    data['event_hooks'] = client.get_all("/api/v1/eventHooks")

    print("  → Inline Hooks...")
    data['inline_hooks'] = client.get_all("/api/v1/inlineHooks")

    print("  → Log Streams...")
    data['log_streams'] = client.get_all("/api/v1/logStreams")

    print("  → ThreatInsight Configuration...")
    data['threat_insight'] = client.get_one("/api/v1/threats/configuration")

    print("  → App Assignments (parallel — this may take a while)...")
    group_map = {g['id']: g.get('profile', {}).get('name', g['id'])
                 for g in data.get('groups', [])}
    data['app_assignments'] = client.fetch_app_assignments_parallel(data['apps'], group_map)

    print("  ✓ Data collection complete!\n")
    return data


# ─── EXCEL BUILDER ────────────────────────────────────────────────────────────

# Color palette — Modern Professional (Charcoal & Teal)
DARK_BLUE   = "263238"   # Charcoal — main titles
MED_BLUE    = "37474F"   # Dark slate — section headers
LIGHT_BLUE  = "ECEFF1"   # Silver-blue — light fills
ACCENT      = "00838F"   # Teal — accent
GREEN       = "00695C"   # Dark teal-green — status text
LIGHT_GREEN = "E0F2F1"   # Pale teal-green — status bg
ORANGE      = "BF360C"   # Burnt orange — warnings
LIGHT_ORANGE= "FBE9E7"   # Pale orange — warning bg
PURPLE      = "283593"   # Deep indigo — groups/policies
LIGHT_PURPLE= "E8EAF6"   # Pale indigo
RED         = "B71C1C"   # Dark red — errors
LIGHT_RED   = "FFEBEE"   # Pale red — error bg
GREY_LIGHT  = "F5F5F5"   # Light grey — alternating rows
GREY_MED    = "CFD8DC"   # Blue-grey — borders/fills
WHITE       = "FFFFFF"
YELLOW      = "FFF8E1"   # Warm pale yellow — manual fields


def hfill(color): return PatternFill("solid", fgColor=color)
def hfont(bold=False, color="000000", size=11, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Arial")
def halign(h="left", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def hborder(style="thin"):
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)

def set_col_width(ws, col, width):
    ws.column_dimensions[get_column_letter(col)].width = width

def header_row(ws, row, cols, fill_color, font_color=WHITE, size=11):
    for c, val in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.fill = hfill(fill_color)
        cell.font = hfont(bold=True, color=font_color, size=size)
        cell.alignment = halign("center")
        cell.border = hborder()

def data_row(ws, row, values, alt=False, color_map=None):
    bg = GREY_LIGHT if alt else WHITE
    for c, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=val)
        cell.fill = hfill(bg)
        cell.font = hfont(size=10)
        cell.alignment = halign(wrap=True)
        cell.border = hborder("hair")

def add_autofilter(ws, header_row_num, col_count, data_count):
    last_col = get_column_letter(col_count)
    end_row = header_row_num + max(data_count, 1)
    ws.auto_filter.ref = f"A{header_row_num}:{last_col}{end_row}"


def build_summary_sheet(wb, data, org):
    ws = wb.active
    ws.title = "📊 Summary"
    ws.sheet_view.showGridLines = False

    # Title block
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    report_org = (org.get('companyName') or 'Okta Tenant').upper()
    c.value = f"🔐  OKTA INVENTORY REPORT — {report_org}"
    c.fill = hfill(DARK_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=16)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 40

    ws.merge_cells("A2:H2")
    c = ws["A2"]
    c.value = f"Generated: {datetime.now().strftime('%B %d, %Y %H:%M')}    |    Org: {org.get('companyName', 'Okta Tenant')}    |    Domain: {org.get('subdomain', '')}.okta.com"
    c.fill = hfill(MED_BLUE)
    c.font = hfont(color=WHITE, size=10)
    c.alignment = halign("center")
    ws.row_dimensions[2].height = 22

    # Spacer
    ws.row_dimensions[3].height = 12

    # ── METRIC CARDS ──
    apps = data.get('apps', [])
    users_active = data.get('users', [])
    users_all = (users_active + data.get('users_staged',[]) +
                 data.get('users_deprovisioned',[]) + data.get('users_suspended',[]) +
                 data.get('users_locked',[]) + data.get('users_password_expired',[]))

    app_active = sum(1 for a in apps if a.get('status') == 'ACTIVE')
    app_inactive = sum(1 for a in apps if a.get('status') != 'ACTIVE')

    saml_apps = sum(1 for a in apps if a.get('signOnMode','').startswith('SAML'))
    oidc_apps = sum(1 for a in apps if a.get('signOnMode','').startswith('OPENID'))
    bookmark_apps = sum(1 for a in apps if a.get('signOnMode','') == 'BOOKMARK')

    oag_apps = sum(1 for a in apps if 'oag' in a.get('label','').lower() or
                   'access gateway' in a.get('label','').lower())

    groups = data.get('groups', [])
    idps = data.get('idps', [])
    agents = data.get('agents', [])
    features = data.get('features', [])
    features_enabled = sum(1 for f in features if f.get('status') == 'ENABLED')
    trusted_origins = data.get('trusted_origins', [])
    tokens = data.get('tokens', [])

    cards = [
        ("APPLICATIONS", [
            ("Total Apps", len(apps)),
            ("Active Apps", app_active),
            ("Inactive Apps", app_inactive),
            ("SAML Apps", saml_apps),
            ("OIDC Apps", oidc_apps),
            ("Bookmark Apps", bookmark_apps),
            ("OAG-Protected Apps", oag_apps),
        ], MED_BLUE, LIGHT_BLUE),
        ("USERS", [
            ("Total Users", len(users_all)),
            ("Active Users", len(users_active)),
            ("Staged Users", len(data.get('users_staged', []))),
            ("Suspended Users", len(data.get('users_suspended', []))),
            ("Locked Out", len(data.get('users_locked', []))),
            ("Password Expired", len(data.get('users_password_expired', []))),
            ("Deprovisioned", len(data.get('users_deprovisioned', []))),
        ], GREEN, LIGHT_GREEN),
        ("DIRECTORIES & IDENTITY", [
            ("Total Groups", len(groups)),
            ("Identity Providers (IdPs)", len(idps)),
            ("Directory Agent Pools", len(agents)),
            ("Trusted Origins", len(trusted_origins)),
            ("Active API Tokens", len(tokens)),
            ("Features Enabled", features_enabled),
            ("Total Features", len(features)),
        ], PURPLE, LIGHT_PURPLE),
    ]

    start_row = 4
    for col_offset, (title, items, dark, light) in enumerate(cards):
        col = col_offset * 3 + 1
        ws.merge_cells(start_row=start_row, start_column=col,
                       end_row=start_row, end_column=col+2)
        c = ws.cell(row=start_row, column=col, value=title)
        c.fill = hfill(dark)
        c.font = hfont(bold=True, color=WHITE, size=11)
        c.alignment = halign("center")
        ws.row_dimensions[start_row].height = 28

        for i, (label, val) in enumerate(items):
            r = start_row + 1 + i
            ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col+1)
            lc = ws.cell(row=r, column=col, value=label)
            lc.fill = hfill(light)
            lc.font = hfont(size=10)
            lc.alignment = halign("left")
            lc.border = hborder("hair")

            vc = ws.cell(row=r, column=col+2, value=val)
            vc.fill = hfill(WHITE)
            vc.font = hfont(bold=True, size=11)
            vc.alignment = halign("center")
            vc.border = hborder("hair")
            ws.row_dimensions[r].height = 22

    # ── POLICIES SECTION ──
    pol_row = start_row + 10
    ws.merge_cells(start_row=pol_row, start_column=1, end_row=pol_row, end_column=8)
    c = ws.cell(row=pol_row, column=1, value="POLICIES & CONFIGURATION")
    c.fill = hfill(ORANGE)
    c.font = hfont(bold=True, color=WHITE, size=11)
    c.alignment = halign("center")
    ws.row_dimensions[pol_row].height = 28

    pol_items = [
        ("Authentication Policies", len(data.get('auth_policies', []))),
        ("MFA Enrollment Policies", len(data.get('mfa_policies', []))),
        ("Password Policies", len(data.get('password_policies', []))),
        ("Sign-On Policies", len(data.get('signon_policies', []))),
    ]
    for i, (label, val) in enumerate(pol_items):
        r = pol_row + 1 + i
        col = (i % 2) * 4 + 1
        ws.merge_cells(start_row=r, start_column=col, end_row=r, end_column=col+2)
        lc = ws.cell(row=r, column=col, value=label)
        lc.fill = hfill(LIGHT_ORANGE)
        lc.font = hfont(size=10)
        lc.alignment = halign("left")
        lc.border = hborder("hair")

        vc = ws.cell(row=r, column=col+3, value=val)
        vc.fill = hfill(WHITE)
        vc.font = hfont(bold=True, size=11)
        vc.alignment = halign("center")
        vc.border = hborder("hair")
        ws.row_dimensions[r].height = 22

    for col, w in enumerate([18, 18, 12, 18, 18, 12, 18, 18], 1):
        set_col_width(ws, col, w)


def build_apps_sheet(wb, apps, app_assignments=None):
    ws = wb.create_sheet("📱 Applications")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:K1")
    c = ws["A1"]
    c.value = f"APPLICATIONS ({len(apps)} total)"
    c.fill = hfill(MED_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["App Name", "Status", "Sign-On Mode", "App Type", "Created",
            "Last Updated", "Assigned Users", "Assigned Groups", "Provisioning", "Label", "App ID"]
    header_row(ws, 2, cols, MED_BLUE)

    widths = [35, 12, 18, 20, 14, 14, 16, 35, 30, 30, 25]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)
    ws.row_dimensions[2].height = 22

    assignments = app_assignments or {}

    for i, app in enumerate(apps):
        r = i + 3
        status = app.get('status', '')
        sign_on = app.get('signOnMode', '')
        created = app.get('created', '')[:10] if app.get('created') else ''
        updated = app.get('lastUpdated', '')[:10] if app.get('lastUpdated') else ''

        asgn = assignments.get(app.get('id', ''), {})
        user_count = asgn.get('user_count', '')
        group_names = ', '.join(asgn.get('group_names', []))
        provisioning = ', '.join(app.get('features', []))

        values = [
            app.get('label', ''),
            status,
            sign_on,
            app.get('name', ''),
            created,
            updated,
            user_count,
            group_names,
            provisioning,
            app.get('label', ''),
            app.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=2)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

        soc = ws.cell(row=r, column=3)
        if 'SAML' in sign_on:
            soc.fill = hfill(LIGHT_BLUE)
        elif 'OPENID' in sign_on:
            soc.fill = hfill(LIGHT_PURPLE)
        elif sign_on == 'BOOKMARK':
            soc.fill = hfill(YELLOW)

    add_autofilter(ws, 2, len(cols), len(apps))


def build_users_sheet(wb, data):
    all_users = []
    for status_key, label in [
        ('users', 'ACTIVE'), ('users_staged', 'STAGED'),
        ('users_suspended', 'SUSPENDED'), ('users_locked', 'LOCKED_OUT'),
        ('users_password_expired', 'PASSWORD_EXPIRED'),
        ('users_deprovisioned', 'DEPROVISIONED')
    ]:
        for u in data.get(status_key, []):
            u['_status_label'] = label
            all_users.append(u)

    ws = wb.create_sheet("👥 Users")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"USERS ({len(all_users)} total)"
    c.fill = hfill(GREEN)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["First Name", "Last Name", "Email / Login", "Status",
            "Created", "Last Login", "Last Updated", "User ID"]
    header_row(ws, 2, cols, GREEN)

    widths = [18, 18, 32, 18, 14, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    status_colors = {
        'ACTIVE': (LIGHT_GREEN, GREEN),
        'STAGED': (YELLOW, "856404"),
        'SUSPENDED': (LIGHT_ORANGE, ORANGE),
        'LOCKED_OUT': (LIGHT_RED, RED),
        'PASSWORD_EXPIRED': ("FFE0B2", "E65100"),
        'DEPROVISIONED': (GREY_MED, "616161"),
    }

    for i, u in enumerate(all_users):
        r = i + 3
        profile = u.get('profile', {})
        status = u.get('_status_label', u.get('status', ''))
        created = u.get('created', '')[:10] if u.get('created') else ''
        last_login = u.get('lastLogin', '')[:10] if u.get('lastLogin') else 'Never'
        updated = u.get('lastUpdated', '')[:10] if u.get('lastUpdated') else ''

        values = [
            profile.get('firstName', ''),
            profile.get('lastName', ''),
            profile.get('email', profile.get('login', '')),
            status,
            created,
            last_login,
            updated,
            u.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=4)
        bg, fg = status_colors.get(status, (WHITE, "000000"))
        sc.fill = hfill(bg)
        sc.font = hfont(color=fg, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(all_users))


def build_groups_sheet(wb, groups):
    ws = wb.create_sheet("🗂️ Groups")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"GROUPS & DIRECTORIES ({len(groups)} total)"
    c.fill = hfill(PURPLE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Group Name", "Type", "Description", "Member Count", "Created", "Group ID"]
    header_row(ws, 2, cols, PURPLE)
    widths = [35, 20, 40, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, g in enumerate(groups):
        r = i + 3
        profile = g.get('profile', {})
        gtype = g.get('type', '')
        created = g.get('created', '')[:10] if g.get('created') else ''

        values = [
            profile.get('name', ''),
            gtype,
            profile.get('description', ''),
            '',
            created,
            g.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        tc = ws.cell(row=r, column=2)
        if gtype == 'OKTA_GROUP':
            tc.fill = hfill(LIGHT_PURPLE)
        elif gtype == 'APP_GROUP':
            tc.fill = hfill(LIGHT_BLUE)
        elif gtype == 'BUILT_IN':
            tc.fill = hfill(YELLOW)

    add_autofilter(ws, 2, len(cols), len(groups))


def build_idp_sheet(wb, idps):
    ws = wb.create_sheet("🔗 Identity Providers")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = f"IDENTITY PROVIDERS / DIRECTORIES ({len(idps)} total)"
    c.fill = hfill(ORANGE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Name", "Type", "Status", "Protocol", "Created", "Last Updated", "IDP ID"]
    header_row(ws, 2, cols, ORANGE)
    widths = [30, 20, 12, 15, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, idp in enumerate(idps):
        r = i + 3
        status = idp.get('status', '')
        created = idp.get('created', '')[:10] if idp.get('created') else ''
        updated = idp.get('lastUpdated', '')[:10] if idp.get('lastUpdated') else ''
        protocol = idp.get('protocol', {}).get('type', '')

        values = [
            idp.get('name', ''),
            idp.get('type', ''),
            status,
            protocol,
            created,
            updated,
            idp.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=3)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(idps))


def build_features_sheet(wb, features):
    ws = wb.create_sheet("⚙️ Features")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = f"OKTA FEATURES ({len(features)} total)"
    c.fill = hfill(DARK_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Feature Name", "Status", "Stage", "Type", "Description"]
    header_row(ws, 2, cols, DARK_BLUE)
    widths = [40, 14, 14, 20, 55]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, f in enumerate(features):
        r = i + 3
        status = f.get('status', '')
        values = [
            f.get('name', ''),
            status,
            f.get('stage', {}).get('value', '') if isinstance(f.get('stage'), dict) else f.get('stage', ''),
            f.get('type', ''),
            f.get('description', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=2)
        if status == 'ENABLED':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(GREY_MED)
            sc.font = hfont(color="616161", bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(features))


def build_policies_sheet(wb, data):
    ws = wb.create_sheet("🛡️ Policies")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    all_policies = []
    for key, label in [
        ('auth_policies', 'Authentication'),
        ('mfa_policies', 'MFA Enrollment'),
        ('password_policies', 'Password'),
        ('signon_policies', 'Sign-On'),
    ]:
        for p in data.get(key, []):
            p['_policy_type_label'] = label
            all_policies.append(p)

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = f"POLICIES ({len(all_policies)} total)"
    c.fill = hfill(PURPLE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Policy Name", "Policy Type", "Status", "Priority", "Created", "Last Updated", "Policy ID"]
    header_row(ws, 2, cols, PURPLE)
    widths = [40, 20, 12, 10, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    type_colors = {
        'Authentication': LIGHT_BLUE,
        'MFA Enrollment': LIGHT_GREEN,
        'Password': LIGHT_ORANGE,
        'Sign-On': LIGHT_PURPLE,
    }

    for i, p in enumerate(all_policies):
        r = i + 3
        status = p.get('status', '')
        ptype = p.get('_policy_type_label', '')
        created = p.get('created', '')[:10] if p.get('created') else ''
        updated = p.get('lastUpdated', '')[:10] if p.get('lastUpdated') else ''

        values = [
            p.get('name', ''),
            ptype,
            status,
            p.get('priority', ''),
            created,
            updated,
            p.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        tc = ws.cell(row=r, column=2)
        tc.fill = hfill(type_colors.get(ptype, WHITE))

        sc = ws.cell(row=r, column=3)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(all_policies))


def build_oag_sheet(wb, apps):
    oag_apps = [a for a in apps if
                'oag' in a.get('label','').lower() or
                'access gateway' in a.get('label','').lower() or
                a.get('signOnMode','') in ['AUTO_LOGIN', 'BOOKMARK'] and
                any(k in a.get('label','').lower() for k in ['portal','vendor','internal','corp'])]

    ws = wb.create_sheet("🔒 OAG Applications")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"OKTA ACCESS GATEWAY (OAG) APPLICATIONS ({len(oag_apps)} identified)"
    c.fill = hfill(RED)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:H2")
    nc = ws["A2"]
    nc.value = "ℹ️  OAG apps are identified by label keywords. Please review and update manually if needed."
    nc.fill = hfill(LIGHT_RED)
    nc.font = hfont(size=9, italic=True, color=RED)
    nc.alignment = halign("center")

    cols = ["App Name", "Status", "Sign-On Mode", "Created", "Last Updated", "App ID", "Backend URL", "Notes"]
    header_row(ws, 3, cols, RED)
    widths = [35, 12, 18, 14, 14, 28, 35, 25]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, app in enumerate(oag_apps):
        r = i + 4
        status = app.get('status', '')
        created = app.get('created', '')[:10] if app.get('created') else ''
        updated = app.get('lastUpdated', '')[:10] if app.get('lastUpdated') else ''

        settings = app.get('settings', {}).get('app', {})
        backend = settings.get('loginUrl', settings.get('url', ''))

        values = [
            app.get('label', ''),
            status,
            app.get('signOnMode', ''),
            created,
            updated,
            app.get('id', ''),
            backend,
            '',
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=2)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

    if not oag_apps:
        ws.cell(row=4, column=1, value="No OAG applications automatically identified. Please populate manually.")

    add_autofilter(ws, 3, len(cols), len(oag_apps))


def build_zones_sheet(wb, zones):
    ws = wb.create_sheet("🌐 Network Zones")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"NETWORK ZONES ({len(zones)} total)"
    c.fill = hfill(ACCENT)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Zone Name", "Type", "Status", "Gateways", "Proxies", "Created", "Last Updated", "Zone ID"]
    header_row(ws, 2, cols, ACCENT)
    widths = [30, 14, 12, 40, 30, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, zone in enumerate(zones):
        r = i + 3
        status = zone.get('status', '')
        created = zone.get('created', '')[:10] if zone.get('created') else ''
        updated = zone.get('lastUpdated', '')[:10] if zone.get('lastUpdated') else ''
        gateways = ', '.join(
            f"{g.get('type','')}: {g.get('value','')}" for g in (zone.get('gateways') or [])
        )
        proxies = ', '.join(
            f"{p.get('type','')}: {p.get('value','')}" for p in (zone.get('proxies') or [])
        )

        values = [
            zone.get('name', ''),
            zone.get('type', ''),
            status,
            gateways,
            proxies,
            created,
            updated,
            zone.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=3)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(zones))


def build_auth_servers_sheet(wb, auth_servers):
    ws = wb.create_sheet("🔑 Auth Servers")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"AUTHORIZATION SERVERS ({len(auth_servers)} total)"
    c.fill = hfill(PURPLE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Name", "Description", "Audiences", "Issuer", "Status", "Created", "Last Updated", "Server ID"]
    header_row(ws, 2, cols, PURPLE)
    widths = [30, 35, 35, 45, 12, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, srv in enumerate(auth_servers):
        r = i + 3
        status = srv.get('status', '')
        created = srv.get('created', '')[:10] if srv.get('created') else ''
        updated = srv.get('lastUpdated', '')[:10] if srv.get('lastUpdated') else ''
        audiences = ', '.join(srv.get('audiences') or [])

        values = [
            srv.get('name', ''),
            srv.get('description', ''),
            audiences,
            srv.get('issuer', ''),
            status,
            created,
            updated,
            srv.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=5)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(auth_servers))


def build_authenticators_sheet(wb, authenticators):
    ws = wb.create_sheet("📲 Authenticators")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = f"AUTHENTICATORS ({len(authenticators)} total)"
    c.fill = hfill(GREEN)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Name", "Key", "Type", "Status", "Created", "Last Updated", "Authenticator ID"]
    header_row(ws, 2, cols, GREEN)
    widths = [30, 25, 20, 12, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, auth in enumerate(authenticators):
        r = i + 3
        status = auth.get('status', '')
        created = auth.get('created', '')[:10] if auth.get('created') else ''
        updated = auth.get('lastUpdated', '')[:10] if auth.get('lastUpdated') else ''

        values = [
            auth.get('name', ''),
            auth.get('key', ''),
            auth.get('type', ''),
            status,
            created,
            updated,
            auth.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=4)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(GREY_MED)
            sc.font = hfont(color="616161", bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(authenticators))


def build_group_rules_sheet(wb, group_rules):
    ws = wb.create_sheet("📏 Group Rules")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = f"GROUP RULES ({len(group_rules)} total)"
    c.fill = hfill(MED_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Rule Name", "Status", "Expression", "Target Groups", "Rule ID"]
    header_row(ws, 2, cols, MED_BLUE)
    widths = [35, 12, 50, 40, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, rule in enumerate(group_rules):
        r = i + 3
        status = rule.get('status', '')
        expression = (rule.get('conditions') or {}).get('expression', {}).get('value', '')
        target_groups = ', '.join(
            (rule.get('actions') or {}).get('assignUserToGroups', {}).get('groupIds') or []
        )

        values = [
            rule.get('name', ''),
            status,
            expression,
            target_groups,
            rule.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=2)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(GREY_MED)
            sc.font = hfont(color="616161", bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(group_rules))


def build_hooks_sheet(wb, event_hooks, inline_hooks):
    all_hooks = (
        [dict(h, _hook_type='Event Hook') for h in event_hooks] +
        [dict(h, _hook_type='Inline Hook') for h in inline_hooks]
    )

    ws = wb.create_sheet("🪝 Hooks")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"EVENT & INLINE HOOKS ({len(all_hooks)} total)"
    c.fill = hfill(ORANGE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Name", "Hook Type", "Status", "Verification", "Endpoint URL", "Events / Trigger", "Created", "Hook ID"]
    header_row(ws, 2, cols, ORANGE)
    widths = [35, 14, 12, 14, 45, 45, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, hook in enumerate(all_hooks):
        r = i + 3
        status = hook.get('status', '')
        hook_type = hook.get('_hook_type', '')
        verification = hook.get('verificationStatus', 'N/A') if hook_type == 'Event Hook' else 'N/A'
        endpoint = hook.get('channel', {}).get('config', {}).get('uri', '')
        created = hook.get('created', '')[:10] if hook.get('created') else ''

        if hook_type == 'Event Hook':
            events = ', '.join((hook.get('events') or {}).get('items') or [])
        else:
            events = hook.get('type', '')

        values = [
            hook.get('name', ''),
            hook_type,
            status,
            verification,
            endpoint,
            events,
            created,
            hook.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=3)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

        tc = ws.cell(row=r, column=2)
        if hook_type == 'Event Hook':
            tc.fill = hfill(LIGHT_ORANGE)
        else:
            tc.fill = hfill(LIGHT_PURPLE)

    add_autofilter(ws, 2, len(cols), len(all_hooks))


def build_log_streams_sheet(wb, log_streams):
    ws = wb.create_sheet("📡 Log Streams")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"LOG STREAMS ({len(log_streams)} total)"
    c.fill = hfill(DARK_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Name", "Type", "Status", "Created", "Last Updated", "Stream ID"]
    header_row(ws, 2, cols, DARK_BLUE)
    widths = [35, 30, 12, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    for i, stream in enumerate(log_streams):
        r = i + 3
        status = stream.get('status', '')
        created = stream.get('created', '')[:10] if stream.get('created') else ''
        updated = stream.get('lastUpdated', '')[:10] if stream.get('lastUpdated') else ''

        values = [
            stream.get('name', ''),
            stream.get('type', ''),
            status,
            created,
            updated,
            stream.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        sc = ws.cell(row=r, column=3)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(LIGHT_RED)
            sc.font = hfont(color=RED, bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(log_streams))


def build_policy_rules_sheet(wb, data):
    ws = wb.create_sheet("📜 Policy Rules")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    policy_lookup = {}
    for key, label in [
        ('auth_policies', 'Authentication'),
        ('mfa_policies', 'MFA Enrollment'),
        ('password_policies', 'Password'),
        ('signon_policies', 'Sign-On'),
    ]:
        for p in data.get(key, []):
            policy_lookup[p['id']] = (p.get('name', ''), label)

    policy_rules = data.get('policy_rules', {})
    all_rows = []
    for pol_id, rules in policy_rules.items():
        pol_name, pol_type = policy_lookup.get(pol_id, (pol_id, ''))
        for rule in rules:
            all_rows.append((pol_name, pol_type, rule))

    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = f"POLICY RULES ({len(all_rows)} total)"
    c.fill = hfill(PURPLE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    cols = ["Policy Name", "Policy Type", "Rule Name", "Priority", "Status", "Access", "Require MFA", "Rule ID"]
    header_row(ws, 2, cols, PURPLE)
    widths = [35, 20, 35, 10, 12, 14, 14, 28]
    for i, w in enumerate(widths, 1):
        set_col_width(ws, i, w)

    type_colors = {
        'Authentication': LIGHT_BLUE,
        'MFA Enrollment': LIGHT_GREEN,
        'Password': LIGHT_ORANGE,
        'Sign-On': LIGHT_PURPLE,
    }

    for i, (pol_name, pol_type, rule) in enumerate(all_rows):
        r = i + 3
        status = rule.get('status', '')
        signon_actions = rule.get('actions', {}).get('signon', {})
        enroll_actions = rule.get('actions', {}).get('enroll', {})
        access = signon_actions.get('access', enroll_actions.get('self', ''))
        require_mfa = signon_actions.get('requireFactor', '')

        values = [
            pol_name,
            pol_type,
            rule.get('name', ''),
            rule.get('priority', ''),
            status,
            str(access) if access != '' else '',
            str(require_mfa) if require_mfa != '' else '',
            rule.get('id', ''),
        ]
        data_row(ws, r, values, alt=(i % 2 == 0))

        tc = ws.cell(row=r, column=2)
        tc.fill = hfill(type_colors.get(pol_type, WHITE))

        sc = ws.cell(row=r, column=5)
        if status == 'ACTIVE':
            sc.fill = hfill(LIGHT_GREEN)
            sc.font = hfont(color=GREEN, bold=True, size=10)
        else:
            sc.fill = hfill(GREY_MED)
            sc.font = hfont(color="616161", bold=True, size=10)

    add_autofilter(ws, 2, len(cols), len(all_rows))


def build_manual_template_sheet(wb):
    ws = wb.create_sheet("📋 Manual Inventory")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    c = ws["A1"]
    c.value = "MANUAL INVENTORY — Items to Fill In"
    c.fill = hfill(MED_BLUE)
    c.font = hfont(bold=True, color=WHITE, size=13)
    c.alignment = halign("center")
    ws.row_dimensions[1].height = 30

    sections = [
        ("OAG CONFIGURATION", [
            ("OAG Version", "", "e.g. 2024.09.0"),
            ("Number of OAG Instances", "", "e.g. 2"),
            ("OAG Server IPs/Hostnames", "", "e.g. cooaggtwprd01, cooaggtwprd02"),
            ("OAG High Availability", "", "Yes / No"),
            ("OAG Protected Apps Count", "", "Total apps behind OAG"),
        ]),
        ("LIFECYCLE MANAGEMENT", [
            ("Provisioning Connectors Active", "", "e.g. Active Directory, Workday"),
            ("Deprovisioning Automated", "", "Yes / No / Partial"),
            ("Joiner-Mover-Leaver Workflow", "", "Configured / Not Configured"),
            ("HR System Integration", "", "e.g. Workday, SAP"),
            ("Okta Workflows Used", "", "Yes / No — list flows if yes"),
        ]),
        ("MFA & AUTHENTICATION", [
            ("MFA Methods Enabled", "", "e.g. Okta Verify, SMS, TOTP"),
            ("Passwordless Enabled", "", "Yes / No"),
            ("Okta FastPass Used", "", "Yes / No"),
            ("Device Trust Configured", "", "Yes / No"),
            ("Risk-Based Authentication", "", "Yes / No"),
        ]),
        ("DIRECTORIES", [
            ("Active Directory Domains", "", "List all connected AD domains"),
            ("AD Agent Version", "", "e.g. 3.17.0"),
            ("LDAP Directories", "", "List if any"),
            ("HR-Mastered Attributes", "", "List attributes mastered from HR"),
            ("Profile Master Source", "", "e.g. Active Directory / Okta"),
        ]),
        ("GOVERNANCE & COMPLIANCE", [
            ("Okta System Log Retention", "", "Days retained"),
            ("SIEM Integration", "", "e.g. Splunk, Sentinel"),
            ("Admin Roles in Use", "", "List super admin count"),
            ("Break-Glass Accounts", "", "Count and location"),
            ("Last Security Review Date", "", "MM/YYYY"),
        ]),
    ]

    row = 3
    for section_title, items in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        c = ws.cell(row=row, column=1, value=section_title)
        c.fill = hfill(MED_BLUE)
        c.font = hfont(bold=True, color=WHITE, size=11)
        c.alignment = halign("left")
        ws.row_dimensions[row].height = 26
        row += 1

        for label, value, hint in items:
            lc = ws.cell(row=row, column=1, value=label)
            lc.fill = hfill(LIGHT_BLUE)
            lc.font = hfont(size=10)
            lc.border = hborder("hair")
            lc.alignment = halign()

            vc = ws.cell(row=row, column=2, value=value)
            vc.fill = hfill(YELLOW)
            vc.font = hfont(size=10, color=MED_BLUE)
            vc.border = hborder("hair")
            vc.alignment = halign(wrap=True)

            hc = ws.cell(row=row, column=3, value=hint)
            ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
            hc.fill = hfill(GREY_LIGHT)
            hc.font = hfont(size=9, italic=True, color="757575")
            hc.border = hborder("hair")
            hc.alignment = halign(wrap=True)

            ws.row_dimensions[row].height = 22
            row += 1

        row += 1

    set_col_width(ws, 1, 35)
    set_col_width(ws, 2, 35)
    set_col_width(ws, 3, 30)
    set_col_width(ws, 4, 20)

    ws.merge_cells(start_row=row+1, start_column=1, end_row=row+1, end_column=4)
    lc = ws.cell(row=row+1, column=1, value="🟡 Yellow cells = Fill in manually     🔵 Blue cells = Labels     ℹ️ Grey cells = Hints")
    lc.fill = hfill(GREY_LIGHT)
    lc.font = hfont(size=9, italic=True)
    lc.alignment = halign("center")


def build_excel(data, output_path):
    print("📊 Building Excel report...")
    wb = openpyxl.Workbook()
    org = data.get('org', {})

    build_summary_sheet(wb, data, org)
    build_apps_sheet(wb, data.get('apps', []), data.get('app_assignments', {}))
    build_users_sheet(wb, data)
    build_groups_sheet(wb, data.get('groups', []))
    build_idp_sheet(wb, data.get('idps', []))
    build_features_sheet(wb, data.get('features', []))
    build_policies_sheet(wb, data)
    build_policy_rules_sheet(wb, data)
    build_oag_sheet(wb, data.get('apps', []))
    build_zones_sheet(wb, data.get('zones', []))
    build_auth_servers_sheet(wb, data.get('auth_servers', []))
    build_authenticators_sheet(wb, data.get('authenticators', []))
    build_group_rules_sheet(wb, data.get('group_rules', []))
    build_hooks_sheet(wb, data.get('event_hooks', []), data.get('inline_hooks', []))
    build_log_streams_sheet(wb, data.get('log_streams', []))
    build_manual_template_sheet(wb)

    wb.save(output_path)
    if isinstance(output_path, str):
        print(f"  ✓ Excel saved: {output_path}")


# ─── DEMO MODE (no API token) ─────────────────────────────────────────────────

def generate_demo_data():
    import random

    sign_on_modes = ['SAML_2_0','OPENID_CONNECT','BOOKMARK','AUTO_LOGIN','SAML_2_0','SAML_2_0']
    app_names = [
        'Vendor Portal Prod','Vendor Portal STG','Workday','ServiceNow','Salesforce',
        'SAP Concur','Office 365','GitHub Enterprise','Jira','Confluence',
        'Tableau','Snowflake','AWS SSO','Azure Portal','Google Workspace',
        'Zoom','Slack','DocuSign','Box','PeopleSoft HCM',
        'Acme Intranet','IT Helpdesk Portal','Finance Portal','HR Self Service',
        'Benefits Portal','Payroll System','Learning Management System',
        'Expense Management','Procurement Portal','Security Dashboard'
    ]

    apps = []
    for i, name in enumerate(app_names):
        apps.append({
            'id': f'0oa{i:015d}',
            'label': name,
            'name': name.lower().replace(' ', '_'),
            'status': 'ACTIVE' if i < 26 else 'INACTIVE',
            'signOnMode': sign_on_modes[i % len(sign_on_modes)],
            'features': random.choice([[], ['PUSH_NEW_USERS'], ['PUSH_NEW_USERS', 'PUSH_PROFILE_UPDATES']]),
            'created': f'202{random.randint(0,4)}-{random.randint(1,12):02d}-{random.randint(1,28):02d}T00:00:00.000Z',
            'lastUpdated': f'2026-{random.randint(1,9):02d}-{random.randint(1,28):02d}T00:00:00.000Z',
        })

    first_names = ['James','Sarah','Michael','Emily','David','Jessica','Robert','Ashley','John','Jennifer']
    last_names = ['Smith','Johnson','Williams','Brown','Jones','Miller','Davis','Garcia','Wilson','Taylor']

    users = []
    for i in range(150):
        fn = first_names[i % len(first_names)]
        ln = last_names[i % len(last_names)]
        users.append({
            'id': f'00u{i:015d}',
            'profile': {'firstName': fn, 'lastName': ln,
                        'email': f'{fn.lower()}.{ln.lower()}{i}@acmecorp.com',
                        'login': f'{fn.lower()}.{ln.lower()}{i}@acmecorp.com'},
            'status': 'ACTIVE',
            'created': f'202{random.randint(0,4)}-{random.randint(1,12):02d}-01T00:00:00.000Z',
            'lastLogin': f'2026-{random.randint(8,9):02d}-{random.randint(1,28):02d}T00:00:00.000Z',
            'lastUpdated': f'2026-{random.randint(1,9):02d}-01T00:00:00.000Z',
        })

    groups = [
        {'id': f'00g{i:015d}', 'type': t, 'created': '2023-01-01T00:00:00.000Z',
         'profile': {'name': n, 'description': d}}
        for i, (n, t, d) in enumerate([
            ('Everyone', 'BUILT_IN', 'All users in org'),
            ('Acme-Employees', 'OKTA_GROUP', 'All Acme employees'),
            ('Acme-Contractors', 'OKTA_GROUP', 'External contractors'),
            ('Acme-IT-Admins', 'OKTA_GROUP', 'IT Administrators'),
            ('Acme-Finance', 'OKTA_GROUP', 'Finance department users'),
            ('Acme-HR', 'OKTA_GROUP', 'Human Resources'),
            ('Acme-Vendors', 'OKTA_GROUP', 'External vendor users'),
            ('Vendor-Portal-Users', 'APP_GROUP', 'Users with Vendor Portal access'),
            ('Workday-Users', 'APP_GROUP', 'Users with Workday access'),
            ('ServiceNow-Users', 'APP_GROUP', 'Users with ServiceNow access'),
        ])
    ]

    idps = [
        {'id': 'idp001', 'name': 'Acme Active Directory', 'type': 'ACTIVE_DIRECTORY',
         'status': 'ACTIVE', 'protocol': {'type': 'LDAP'},
         'created': '2021-06-01T00:00:00.000Z', 'lastUpdated': '2026-01-15T00:00:00.000Z'},
        {'id': 'idp002', 'name': 'Azure AD Federation', 'type': 'MICROSOFT',
         'status': 'ACTIVE', 'protocol': {'type': 'OIDC'},
         'created': '2022-03-01T00:00:00.000Z', 'lastUpdated': '2026-02-10T00:00:00.000Z'},
    ]

    features = [
        {'name': 'Okta Access Gateway', 'status': 'ENABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Protects on-prem web applications'},
        {'name': 'Okta Workflows', 'status': 'ENABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'No-code automation for identity'},
        {'name': 'Device Trust', 'status': 'ENABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Enforce device compliance for access'},
        {'name': 'Passwordless Authentication', 'status': 'ENABLED', 'stage': {'value': 'BETA'}, 'type': 'self-service', 'description': 'FastPass and FIDO2 WebAuthn'},
        {'name': 'Risk-Based Authentication', 'status': 'DISABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Adaptive MFA based on risk signals'},
        {'name': 'Identity Governance', 'status': 'DISABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Access certifications and entitlements'},
        {'name': 'Okta Identity Engine', 'status': 'ENABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Next-gen identity platform'},
        {'name': 'Universal Directory', 'status': 'ENABLED', 'stage': {'value': 'GA'}, 'type': 'self-service', 'description': 'Flexible user store'},
    ]

    auth_policies = [
        {'id': f'pol{i:010d}', 'name': n, 'status': 'ACTIVE', 'priority': i+1,
         'created': '2023-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'}
        for i, n in enumerate(['Default Policy','Vendor Portal Policy','High Security Apps',
                                'MFA Required Policy','Passwordless Policy'])
    ]

    zones = [
        {'id': 'nzo001', 'name': 'Corporate Network', 'type': 'IP', 'status': 'ACTIVE',
         'gateways': [{'type': 'RANGE', 'value': '10.0.0.0-10.255.255.255'}],
         'proxies': [],
         'created': '2021-06-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
        {'id': 'nzo002', 'name': 'Vendor Access Zone', 'type': 'IP', 'status': 'ACTIVE',
         'gateways': [{'type': 'CIDR', 'value': '203.0.113.0/24'}],
         'proxies': [{'type': 'CIDR', 'value': '198.51.100.0/24'}],
         'created': '2022-03-01T00:00:00.000Z', 'lastUpdated': '2026-02-01T00:00:00.000Z'},
        {'id': 'nzo003', 'name': 'BlockList Zone', 'type': 'IP', 'status': 'ACTIVE',
         'gateways': [{'type': 'CIDR', 'value': '192.0.2.0/24'}],
         'proxies': [],
         'created': '2023-01-01T00:00:00.000Z', 'lastUpdated': '2025-11-01T00:00:00.000Z'},
    ]

    auth_servers = [
        {'id': 'aus001', 'name': 'Acme API Gateway', 'description': 'Custom OAuth server for internal APIs',
         'audiences': ['api://acme-internal'], 'issuer': 'https://acme.okta.com/oauth2/aus001',
         'status': 'ACTIVE', 'created': '2022-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
        {'id': 'aus002', 'name': 'Vendor Portal Auth', 'description': 'OAuth server for vendor-facing APIs',
         'audiences': ['api://acme-vendor'], 'issuer': 'https://acme.okta.com/oauth2/aus002',
         'status': 'ACTIVE', 'created': '2022-06-01T00:00:00.000Z', 'lastUpdated': '2025-12-01T00:00:00.000Z'},
        {'id': 'aus003', 'name': 'Mobile Apps Auth', 'description': 'OAuth server for mobile applications',
         'audiences': ['api://acme-mobile'], 'issuer': 'https://acme.okta.com/oauth2/aus003',
         'status': 'INACTIVE', 'created': '2023-03-01T00:00:00.000Z', 'lastUpdated': '2025-06-01T00:00:00.000Z'},
    ]

    authenticators = [
        {'id': 'aut001', 'key': 'okta_verify', 'type': 'APP', 'name': 'Okta Verify',
         'status': 'ACTIVE', 'created': '2021-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
        {'id': 'aut002', 'key': 'google_otp', 'type': 'APP', 'name': 'Google Authenticator',
         'status': 'ACTIVE', 'created': '2021-01-01T00:00:00.000Z', 'lastUpdated': '2025-06-01T00:00:00.000Z'},
        {'id': 'aut003', 'key': 'okta_password', 'type': 'PASSWORD', 'name': 'Password',
         'status': 'ACTIVE', 'created': '2021-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
        {'id': 'aut004', 'key': 'phone_number', 'type': 'PHONE', 'name': 'Phone / SMS',
         'status': 'ACTIVE', 'created': '2021-01-01T00:00:00.000Z', 'lastUpdated': '2025-09-01T00:00:00.000Z'},
        {'id': 'aut005', 'key': 'webauthn', 'type': 'SECURITY_KEY', 'name': 'Security Key or Biometric',
         'status': 'ACTIVE', 'created': '2023-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
    ]

    group_rules = [
        {'id': 'rule001', 'name': 'Assign Finance Users', 'status': 'ACTIVE',
         'conditions': {'expression': {'value': 'user.department == "Finance"', 'type': 'urn:okta:expression:1.0'}},
         'actions': {'assignUserToGroups': {'groupIds': ['00g0000000004000000']}}},
        {'id': 'rule002', 'name': 'Assign Vendor Users', 'status': 'ACTIVE',
         'conditions': {'expression': {'value': 'user.userType == "Contractor"', 'type': 'urn:okta:expression:1.0'}},
         'actions': {'assignUserToGroups': {'groupIds': ['00g0000000006000000']}}},
        {'id': 'rule003', 'name': 'Assign IT Admins', 'status': 'INACTIVE',
         'conditions': {'expression': {'value': 'user.department == "IT" AND user.title contains "Admin"', 'type': 'urn:okta:expression:1.0'}},
         'actions': {'assignUserToGroups': {'groupIds': ['00g0000000003000000']}}},
    ]

    event_hooks = [
        {'id': 'whe001', 'name': 'User Lifecycle Events', 'status': 'ACTIVE',
         'verificationStatus': 'VERIFIED',
         'channel': {'config': {'uri': 'https://automation.acmecorp.com/okta/lifecycle'}},
         'events': {'items': ['user.lifecycle.activate', 'user.lifecycle.deactivate', 'user.lifecycle.suspend']},
         'created': '2023-01-01T00:00:00.000Z'},
        {'id': 'whe002', 'name': 'MFA Events', 'status': 'ACTIVE',
         'verificationStatus': 'VERIFIED',
         'channel': {'config': {'uri': 'https://siem.acmecorp.com/okta/mfa'}},
         'events': {'items': ['user.mfa.factor.activate', 'user.mfa.factor.deactivate']},
         'created': '2023-06-01T00:00:00.000Z'},
    ]

    inline_hooks = [
        {'id': 'cal001', 'name': 'Token Enrichment Hook', 'status': 'ACTIVE',
         'type': 'com.okta.oauth2.tokens.transform',
         'channel': {'config': {'uri': 'https://api.acmecorp.com/okta/token-hook'}},
         'created': '2023-03-01T00:00:00.000Z'},
        {'id': 'cal002', 'name': 'User Import Hook', 'status': 'ACTIVE',
         'type': 'com.okta.import.transform',
         'channel': {'config': {'uri': 'https://api.acmecorp.com/okta/import-hook'}},
         'created': '2023-09-01T00:00:00.000Z'},
    ]

    log_streams = [
        {'id': 'lst001', 'name': 'Splunk SIEM Integration', 'type': 'splunk_cloud_logstreaming',
         'status': 'ACTIVE', 'created': '2022-01-01T00:00:00.000Z', 'lastUpdated': '2026-01-01T00:00:00.000Z'},
        {'id': 'lst002', 'name': 'AWS EventBridge Stream', 'type': 'aws_eventbridge',
         'status': 'ACTIVE', 'created': '2023-06-01T00:00:00.000Z', 'lastUpdated': '2025-12-01T00:00:00.000Z'},
    ]

    threat_insight = {
        'action': 'audit',
        'excludeZones': ['nzo001'],
    }

    policy_rules = {
        'pol0000000000': [
            {'id': 'rul0000000001', 'name': 'Default Rule', 'priority': 1, 'status': 'ACTIVE',
             'actions': {'signon': {'access': 'ALLOW', 'requireFactor': True}}},
        ],
        'pol0000000001': [
            {'id': 'rul0000000002', 'name': 'Vendor MFA Rule', 'priority': 1, 'status': 'ACTIVE',
             'actions': {'signon': {'access': 'ALLOW', 'requireFactor': True}}},
            {'id': 'rul0000000003', 'name': 'Fallback Deny', 'priority': 99, 'status': 'ACTIVE',
             'actions': {'signon': {'access': 'DENY', 'requireFactor': False}}},
        ],
        'pol0000000002': [
            {'id': 'rul0000000004', 'name': 'Require Strong MFA', 'priority': 1, 'status': 'ACTIVE',
             'actions': {'signon': {'access': 'ALLOW', 'requireFactor': True}}},
        ],
        'pol0000000003': [
            {'id': 'rul0000000005', 'name': 'Enforce MFA Enrollment', 'priority': 1, 'status': 'ACTIVE',
             'actions': {'enroll': {'self': 'REQUIRED'}}},
        ],
        'pol0000000004': [
            {'id': 'rul0000000006', 'name': 'Passwordless Rule', 'priority': 1, 'status': 'ACTIVE',
             'actions': {'signon': {'access': 'ALLOW', 'requireFactor': True}}},
        ],
    }

    return {
        'org': {'companyName': 'Acme Corporation', 'subdomain': 'acme'},
        'apps': apps,
        'users': users,
        'users_staged': users[:5],
        'users_suspended': users[:3],
        'users_locked': users[:2],
        'users_password_expired': users[:4],
        'users_deprovisioned': users[:8],
        'groups': groups,
        'idps': idps,
        'agents': [{'id': 'agent001', 'name': 'Acme-AD-Agent-Pool'}],
        'auth_policies': auth_policies,
        'mfa_policies': auth_policies[:2],
        'password_policies': auth_policies[:3],
        'signon_policies': auth_policies[:2],
        'features': features,
        'trusted_origins': [{'id': 'to001', 'origin': 'https://vendors.acmecorp.com'},
                             {'id': 'to002', 'origin': 'https://vendor-stg.acmecorp.com'}],
        'tokens': [{'id': 'tok001', 'name': 'Automation-Token'},
                   {'id': 'tok002', 'name': 'Monitoring-Token'}],
        'app_assignments': {
            f'0oa{i:015d}': {
                'user_count': random.randint(5, 200),
                'group_names': random.sample(
                    ['Acme-Employees', 'Acme-Contractors', 'Acme-IT-Admins',
                     'Acme-Finance', 'Acme-HR', 'Acme-Vendors'], k=random.randint(1, 3)
                ),
            }
            for i in range(len(app_names))
        },
        'zones': zones,
        'auth_servers': auth_servers,
        'authenticators': authenticators,
        'group_rules': group_rules,
        'event_hooks': event_hooks,
        'inline_hooks': inline_hooks,
        'log_streams': log_streams,
        'threat_insight': threat_insight,
        'policy_rules': policy_rules,
    }


# ─── MAIN ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Okta Inventory Script")
    parser.add_argument("--url", help="Okta org URL, e.g. https://acme.okta.com")
    parser.add_argument("--token", help="Okta API token (SSWS)")
    parser.add_argument("--output", default="Okta_Inventory.xlsx", help="Output Excel filename")
    parser.add_argument("--demo", action="store_true", help="Generate demo/template with sample data")
    args = parser.parse_args()

    if args.demo or (not args.url or not args.token):
        print("🔄 Running in DEMO mode with sample data...")
        print("   (Run with --url and --token to pull live Okta data)\n")
        data = generate_demo_data()
    else:
        client = OktaClient(args.url, args.token)
        data = collect_all(client)

    output = args.output
    build_excel(data, output)
    print(f"\n✅ Done! Open: {output}")
    print("\nTo run with live Okta data:")
    print(f"   python okta_inventory.py --url https://YOUR-ORG.okta.com --token YOUR_API_TOKEN")


if __name__ == "__main__":
    main()
